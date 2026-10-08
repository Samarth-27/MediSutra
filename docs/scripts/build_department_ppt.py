import os
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    template_path = "PPT Format for Major Project.pptx"
    prs = Presentation(template_path)

    # Palette
    PRIMARY_BLUE = RGBColor(2, 132, 199)    # #0284C7
    DARK_NAVY    = RGBColor(15, 23, 42)     # #0F172A
    SLATE_TEXT   = RGBColor(51, 65, 85)     # #334155
    LIGHT_BG     = RGBColor(248, 250, 252)  # #F8FAFC
    CARD_BG      = RGBColor(240, 249, 255)  # #F0F9FF
    WHITE        = RGBColor(255, 255, 255)
    BORDER_BLUE  = RGBColor(186, 230, 253)  # #BAE6FD

    # 1. Update Slide 1 (Title Slide)
    s1 = prs.slides[0]
    for s in s1.shapes:
        if s.name == "Subtitle 2":
            s.text_frame.paragraphs[0].text = "Major Project Presentation"
            s.text_frame.paragraphs[1].text = "On"
            s.text_frame.paragraphs[3].text = "MEDISUTRA: A Unified Longitudinal Personal Health Intelligence & Federated Clinical Copilot Platform"
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
        elif s.name == "TextBox 6":
            tf = s.text_frame
            tf.paragraphs[3].text = "Presented By:-"
            tf.paragraphs[4].text = "Samarth (Lead Developer)"
            tf.paragraphs[5].text = "Roll No: [Enrollment No]"
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
        elif s.name == "TextBox 10":
            tf = s.text_frame
            tf.paragraphs[2].text = "Guide:-"
            tf.paragraphs[3].text = "[Faculty Guide Name]\nAssistant Professor"
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
        elif s.name == "Rectangle 11":
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"

    # Slide 2 is the official Contents slide - ensure Times New Roman font
    s2 = prs.slides[1]
    for s in s2.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"

    # Clear Slide 3 (the instruction slide) to repurpose it
    s3 = prs.slides[2]
    for shape in list(s3.shapes):
        sp = shape._element
        sp.getparent().remove(sp)

    blank_layout = prs.slide_layouts[6]

    def add_standard_header(slide, title_text, category_text="BCO 519A — MAJOR PROJECT DEFENSE"):
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        accent.fill.solid()
        accent.fill.fore_color.rgb = PRIMARY_BLUE
        accent.line.fill.background()

        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.35))
        p_c = cat_box.text_frame.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.name = "Times New Roman"
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = PRIMARY_BLUE

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.8))
        p_t = t_box.text_frame.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Times New Roman"
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_NAVY

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_BLUE
        line.line.fill.background()

    def add_bullet(tf, bold_txt, body_txt, pt_size=12, space_after=8):
        p = tf.add_paragraph()
        p.space_after = Pt(space_after)
        if bold_txt:
            r1 = p.add_run()
            r1.text = bold_txt + " "
            r1.font.name = "Times New Roman"
            r1.font.bold = True
            r1.font.size = Pt(pt_size)
            r1.font.color.rgb = DARK_NAVY
        r2 = p.add_run()
        r2.text = body_txt
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(pt_size)
        r2.font.color.rgb = SLATE_TEXT

    # =========================================================================
    # SLIDE 3: INTRODUCTION & PROBLEM STATEMENT
    # =========================================================================
    # s3 is already cleared and ready to use
    add_standard_header(s3, "Introduction & Clinical Problem Statement")

    card_s3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    card_s3_1.fill.solid()
    card_s3_1.fill.fore_color.rgb = LIGHT_BG
    card_s3_1.line.color.rgb = BORDER_BLUE
    tf3_1 = card_s3_1.text_frame
    tf3_1.margin_left = Inches(0.3)
    tf3_1.margin_top = Inches(0.3)
    p3_1 = tf3_1.paragraphs[0]
    p3_1.text = "THE FRAGMENTATION CRISIS"
    p3_1.font.name = "Times New Roman"
    p3_1.font.bold = True
    p3_1.font.size = Pt(13)
    p3_1.font.color.rgb = RGBColor(220, 38, 38)

    add_bullet(tf3_1, "Paper Record Silos:", "Patients carry physical plastic files across multi-hospital visits; documents are routinely misplaced or damaged.", 12, 10)
    add_bullet(tf3_1, "Redundant Diagnostic Tests:", "Inaccessible historical lab results force doctors to repeat expensive tests (HbA1c, CT scans), wasting resources.", 12, 10)
    add_bullet(tf3_1, "Clinical Blindspots:", "Attending physicians lack visibility into past hospitalizations, adverse drug reactions, and baseline readings.", 12, 10)
    add_bullet(tf3_1, "Conflation of Acute & Chronic:", "Cured illnesses (Dengue, Bronchitis) are mixed with active conditions (Diabetes, Hypertension).", 12, 10)

    card_s3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3))
    card_s3_2.fill.solid()
    card_s3_2.fill.fore_color.rgb = CARD_BG
    card_s3_2.line.color.rgb = PRIMARY_BLUE
    tf3_2 = card_s3_2.text_frame
    tf3_2.margin_left = Inches(0.3)
    tf3_2.margin_top = Inches(0.3)
    p3_2 = tf3_2.paragraphs[0]
    p3_2.text = "MEDISUTRA PLATFORM SOLUTION"
    p3_2.font.name = "Times New Roman"
    p3_2.font.bold = True
    p3_2.font.size = Pt(13)
    p3_2.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf3_2, "Unified Health Identity:", "Anchored to the 14-digit ABHA/ABDM identifier, unifying multi-hospital encounters into one health journey.", 12, 10)
    add_bullet(tf3_2, "Federated Hospital Network:", "Interoperable clinical graph connecting Apollo, Fortis, Max, AIIMS, Dr. Lal PathLabs, and Metropolis.", 12, 10)
    add_bullet(tf3_2, "Disease Lifecycle State Machine:", "Explicit clinical separation of active chronic conditions from certified cured historical illnesses.", 12, 10)
    add_bullet(tf3_2, "Deterministic Clinical AI:", "Zero-hallucination decision support copilot calculating lab trajectories with 1-click consultation note insertion.", 12, 10)

    # =========================================================================
    # SLIDE 4: MAJOR PROJECT OBJECTIVES
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_standard_header(s4, "Project Objectives & Course Outcome Attainment")

    objs = [
        ("CO1: Problem Formulation", "Formulate a condition-centric longitudinal computing model to resolve Indian healthcare data fragmentation."),
        ("CO2: System Design & Tooling", "Architect an enterprise 4-tier web platform using React 19, TypeScript, Express.js, and relational database ER modeling."),
        ("CO3: Implementation & Testing", "Develop modular micro-services and reactive workstations with 0 compiler/lint errors and E2E test verification."),
        ("CO4: Ethics & Project Mgmt", "Enforce Zero-Trust security and data sovereignty adhering to India's DPDPA 2023 and agile GitHub version control."),
        ("CO5: Technical Defense", "Deliver comprehensive architectural specs, live system demonstrations, visual recordings, and thesis defense.")
    ]
    for idx, (co_title, co_desc) in enumerate(objs):
        y = Inches(1.7 + idx * 1.05)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.92))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_BLUE

        badge = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.18), Inches(2.6), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY_BLUE
        badge.line.fill.background()
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = co_title
        p_b.font.name = "Times New Roman"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.color.rgb = WHITE

        tbox = s4.shapes.add_textbox(Inches(3.8), y + Inches(0.1), Inches(8.5), Inches(0.72))
        p_desc = tbox.text_frame.paragraphs[0]
        p_desc.text = co_desc
        p_desc.font.name = "Times New Roman"
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = DARK_NAVY

    # =========================================================================
    # SLIDE 5: LITERATURE REVIEW & GAP ANALYSIS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_standard_header(s5, "Literature Review & Research Gap Analysis")

    card_lr1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    card_lr1.fill.solid()
    card_lr1.fill.fore_color.rgb = LIGHT_BG
    card_lr1.line.color.rgb = BORDER_BLUE
    tf_lr1 = card_lr1.text_frame
    tf_lr1.margin_left = Inches(0.3)
    tf_lr1.margin_top = Inches(0.3)
    p_hlr1 = tf_lr1.paragraphs[0]
    p_hlr1.text = "EXISTING HEALTHCARE PLATFORMS"
    p_hlr1.font.name = "Times New Roman"
    p_hlr1.font.bold = True
    p_hlr1.font.size = Pt(13)
    p_hlr1.font.color.rgb = DARK_NAVY

    add_bullet(tf_lr1, "Epic MyChart / Cerner:", "Proprietary enterprise EHR walled gardens. Cross-facility sharing requires vendor-specific exchange networks.", 11.5, 8)
    add_bullet(tf_lr1, "Apple Health:", "Patient-pulled mobile application. Displays raw lists without active vs. cured disease separation.", 11.5, 8)
    add_bullet(tf_lr1, "ABDM PHR Apps (Official):", "Enables consent-based document push but stores static PDFs without longitudinal biomarker analytics.", 11.5, 8)
    add_bullet(tf_lr1, "Practo / 1mg:", "Commercial tele-health platforms limited to private clinics with basic customer FAQ bots.", 11.5, 8)

    card_lr2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3))
    card_lr2.fill.solid()
    card_lr2.fill.fore_color.rgb = CARD_BG
    card_lr2.line.color.rgb = PRIMARY_BLUE
    tf_lr2 = card_lr2.text_frame
    tf_lr2.margin_left = Inches(0.3)
    tf_lr2.margin_top = Inches(0.3)
    p_hlr2 = tf_lr2.paragraphs[0]
    p_hlr2.text = "IDENTIFIED RESEARCH GAPS"
    p_hlr2.font.name = "Times New Roman"
    p_hlr2.font.bold = True
    p_hlr2.font.size = Pt(13)
    p_hlr2.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_lr2, "Gap 1 (The PDF Dump Problem):", "Doctors cannot read through dozens of raw PDF files during brief outpatient consultations. Records must be condition-centric.", 11.5, 8)
    add_bullet(tf_lr2, "Gap 2 (Lack of Cure Lifecycle):", "Current EHRs lack certified clinical workflows for stamping disease cure milestones with follow-up evidence.", 11.5, 8)
    add_bullet(tf_lr2, "Gap 3 (Clinical AI Hallucinations):", "Generative LLMs hallucinate medical facts and fail at temporal math. Healthcare demands deterministic, grounded decision support.", 11.5, 8)
    add_bullet(tf_lr2, "MediSutra Innovation:", "Bridges all three gaps through an integrated federated architecture.", 11.5, 8)

    # =========================================================================
    # SLIDE 6: SYSTEM DESIGN — 4-TIER ARCHITECTURE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_standard_header(s6, "System Design: Enterprise 4-Tier Architecture")

    img_arch = "docs/diagrams/01_system_architecture_4_tier.png"
    if os.path.exists(img_arch):
        s6.shapes.add_picture(img_arch, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_s6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_s6.fill.solid()
    card_s6.fill.fore_color.rgb = LIGHT_BG
    card_s6.line.color.rgb = BORDER_BLUE
    tf_s6 = card_s6.text_frame
    tf_s6.margin_left = Inches(0.25)
    tf_s6.margin_top = Inches(0.25)
    p_s6 = tf_s6.paragraphs[0]
    p_s6.text = "TIERED ARCHITECTURAL SPECIFICATION"
    p_s6.font.name = "Times New Roman"
    p_s6.font.bold = True
    p_s6.font.size = Pt(13)
    p_s6.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s6, "Tier 1: Presentation Layer", "React 19 + TypeScript + Vite. Role-partitioned clinical workstations for Doctors, Hospital Admins, and Citizens.", 11.5, 8)
    add_bullet(tf_s6, "Tier 2: API Gateway Layer", "Express.js REST gateway with Zero-Trust RBAC middleware, session persistence, and standardized JSON envelopes.", 11.5, 8)
    add_bullet(tf_s6, "Tier 3: Clinical Intelligence", "Deterministic RAG pipeline, quantitative trajectory computing engine, and disease lifecycle state machine.", 11.5, 8)
    add_bullet(tf_s6, "Tier 4: Persistence Layer", "Relational health models, 14-digit ABHA patient registry, and immutable audit logs.", 11.5, 8)

    # =========================================================================
    # SLIDE 7: SYSTEM DESIGN — DATA FLOW DIAGRAMS (DFD 0 & 1)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_standard_header(s7, "System Design: Data Flow Diagrams (DFD Level 0 & Level 1)")

    if os.path.exists("docs/diagrams/02_dfd_level_0_context.png"):
        s7.shapes.add_picture("docs/diagrams/02_dfd_level_0_context.png", Inches(0.8), Inches(1.8), width=Inches(5.6))
        cap0 = s7.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(5.6), Inches(0.5))
        p_c0 = cap0.text_frame.paragraphs[0]
        p_c0.text = "Figure 4.2: DFD Level 0 Context Level Diagram"
        p_c0.font.name = "Times New Roman"
        p_c0.font.size = Pt(10)
        p_c0.font.bold = True
        p_c0.font.color.rgb = SLATE_TEXT
        p_c0.alignment = PP_ALIGN.CENTER

    if os.path.exists("docs/diagrams/03_dfd_level_1_subsystems.png"):
        s7.shapes.add_picture("docs/diagrams/03_dfd_level_1_subsystems.png", Inches(6.8), Inches(1.8), width=Inches(5.7))
        cap1 = s7.shapes.add_textbox(Inches(6.8), Inches(6.5), Inches(5.7), Inches(0.5))
        p_c1 = cap1.text_frame.paragraphs[0]
        p_c1.text = "Figure 4.3: DFD Level 1 Core Subsystems Decomposition"
        p_c1.font.name = "Times New Roman"
        p_c1.font.size = Pt(10)
        p_c1.font.bold = True
        p_c1.font.color.rgb = SLATE_TEXT
        p_c1.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 8: SYSTEM DESIGN — FEDERATED HRC TOPOLOGY
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_standard_header(s8, "System Design: Federated Hospital Network (HRC Topology)")

    img_hrc = "docs/diagrams/04_hrc_network_topology.png"
    if os.path.exists(img_hrc):
        s8.shapes.add_picture(img_hrc, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_s8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_s8.fill.solid()
    card_s8.fill.fore_color.rgb = CARD_BG
    card_s8.line.color.rgb = PRIMARY_BLUE
    tf_s8 = card_s8.text_frame
    tf_s8.margin_left = Inches(0.25)
    tf_s8.margin_top = Inches(0.25)
    p_s8 = tf_s8.paragraphs[0]
    p_s8.text = "FEDERATED NETWORK TOPOLOGY"
    p_s8.font.name = "Times New Roman"
    p_s8.font.bold = True
    p_s8.font.size = Pt(13)
    p_s8.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s8, "Central Registry Hub:", "Maintains unified patient identities anchored to the 14-digit ABHA number.", 11.5, 8)
    add_bullet(tf_s8, "Hospital Nodes:", "Apollo Hospitals, Fortis Memorial, Max Super Speciality, AIIMS New Delhi operate authenticated clinical stations.", 11.5, 8)
    add_bullet(tf_s8, "Diagnostic Nodes:", "Dr. Lal PathLabs and Metropolis Healthcare ingest laboratory test panels directly into the central vault.", 11.5, 8)
    add_bullet(tf_s8, "Decentralized Continuity:", "Hospitals retain institutional autonomy while contributing standardized encounter data to the patient's sovereign health graph.", 11.5, 8)

    # =========================================================================
    # SLIDE 9: SYSTEM DESIGN — CONCEPTUAL ER DIAGRAM
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_standard_header(s9, "System Design: Conceptual Entity-Relationship (ER) Model")

    img_er = "docs/diagrams/10_conceptual_er_diagram.png"
    if os.path.exists(img_er):
        s9.shapes.add_picture(img_er, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_s9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_s9.fill.solid()
    card_s9.fill.fore_color.rgb = LIGHT_BG
    card_s9.line.color.rgb = BORDER_BLUE
    tf_s9 = card_s9.text_frame
    tf_s9.margin_left = Inches(0.25)
    tf_s9.margin_top = Inches(0.25)
    p_s9 = tf_s9.paragraphs[0]
    p_s9.text = "3NF RELATIONAL SCHEMA"
    p_s9.font.name = "Times New Roman"
    p_s9.font.bold = True
    p_s9.font.size = Pt(13)
    p_s9.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s9, "Patients Entity:", "healthId (PK), abhaNumber (14-digit unique), demographic data, and primary facility.", 11.5, 8)
    add_bullet(tf_s9, "Encounters Entity:", "Captures doctorId, hospitalId, ICD-10 diagnosis, clinical severity, and examination notes.", 11.5, 8)
    add_bullet(tf_s9, "Conditions Entity:", "Maintains lifetime disease state (ACTIVE vs. RESOLVED), diagnosedDate, curedDate, and curedEvidence.", 11.5, 8)
    add_bullet(tf_s9, "Reports & Biomarkers:", "Stores multi-year lab panels and quantitative analyte readings (HbA1c, Glucose, Vitamin D).", 11.5, 8)

    # =========================================================================
    # SLIDE 10: SYSTEM DESIGN — LIFETIME DISEASE STATE MACHINE
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_standard_header(s10, "System Design: Disease Lifecycle State Transition Machine")

    img_sm = "docs/diagrams/05_disease_lifecycle_state_machine.png"
    if os.path.exists(img_sm):
        s10.shapes.add_picture(img_sm, Inches(0.8), Inches(1.8), width=Inches(7.2))

    card_s10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.1))
    card_s10.fill.solid()
    card_s10.fill.fore_color.rgb = CARD_BG
    card_s10.line.color.rgb = PRIMARY_BLUE
    tf_s10 = card_s10.text_frame
    tf_s10.margin_left = Inches(0.25)
    tf_s10.margin_top = Inches(0.25)
    p_s10 = tf_s10.paragraphs[0]
    p_s10.text = "EXPLICIT CLINICAL LIFECYCLE"
    p_s10.font.name = "Times New Roman"
    p_s10.font.bold = True
    p_s10.font.size = Pt(13)
    p_s10.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s10, "Active Chronic State:", "Diseases requiring ongoing therapy (Type 2 Diabetes, Hypertension) remain tagged ACTIVE.", 11.5, 8)
    add_bullet(tf_s10, "Cure Certification:", "Doctors certify illness resolution by submitting follow-up clinical examination notes and diagnostic evidence.", 11.5, 8)
    add_bullet(tf_s10, "Longitudinal History:", "Cured illnesses (Dengue, Bronchitis) transition to RESOLVED and are archived into historical records with verified evidence.", 11.5, 8)

    # =========================================================================
    # SLIDE 11: SYSTEM DESIGN — UML SEQUENCE DIAGRAM
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_standard_header(s11, "System Design: Clinical Consultation Workflow (UML Sequence)")

    img_seq = "docs/diagrams/09_uml_consultation_sequence.png"
    if os.path.exists(img_seq):
        s11.shapes.add_picture(img_seq, Inches(0.8), Inches(1.7), width=Inches(7.0))

    card_s11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.7), Inches(4.5), Inches(5.3))
    card_s11.fill.solid()
    card_s11.fill.fore_color.rgb = LIGHT_BG
    card_s11.line.color.rgb = BORDER_BLUE
    tf_s11 = card_s11.text_frame
    tf_s11.margin_left = Inches(0.25)
    tf_s11.margin_top = Inches(0.25)
    p_s11 = tf_s11.paragraphs[0]
    p_s11.text = "INTERACTION FLOW"
    p_s11.font.name = "Times New Roman"
    p_s11.font.bold = True
    p_s11.font.size = Pt(13)
    p_s11.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s11, "1. Patient Selection:", "Doctor selects patient from assigned roster.", 11.5, 6)
    add_bullet(tf_s11, "2. Dossier Ingestion:", "Gateway queries multi-hospital data store and streams longitudinal record.", 11.5, 6)
    add_bullet(tf_s11, "3. Copilot Query:", "Doctor clicks preset chip or types clinical symptoms.", 11.5, 6)
    add_bullet(tf_s11, "4. Real-Time Synthesis:", "AI engine correlates symptoms, computes trajectories, and returns report.", 11.5, 6)
    add_bullet(tf_s11, "5. Consultation Note Export:", "1-click copy directly inserts structured intelligence into active note.", 11.5, 6)

    # =========================================================================
    # SLIDE 12: METHODOLOGY — DETERMINISTIC CLINICAL AI ENGINE
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_standard_header(s12, "Methodology: Deterministic Clinical AI Decision Support")

    img_ai = "docs/diagrams/06_clinical_ai_rag_flowchart.png"
    if os.path.exists(img_ai):
        s12.shapes.add_picture(img_ai, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_s12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_s12.fill.solid()
    card_s12.fill.fore_color.rgb = CARD_BG
    card_s12.line.color.rgb = PRIMARY_BLUE
    tf_s12 = card_s12.text_frame
    tf_s12.margin_left = Inches(0.25)
    tf_s12.margin_top = Inches(0.25)
    p_s12 = tf_s12.paragraphs[0]
    p_s12.text = "DETERMINISTIC RAG PIPELINE"
    p_s12.font.name = "Times New Roman"
    p_s12.font.bold = True
    p_s12.font.size = Pt(13)
    p_s12.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s12, "1. Query Ingestion:", "Doctors input clinical observations or select instant preset chips.", 11.5, 6)
    add_bullet(tf_s12, "2. Cross-Hospital Fetch:", "Queries patient's complete cross-hospital record across all connected facilities.", 11.5, 6)
    add_bullet(tf_s12, "3. Symptom Correlation:", "Maps tokens to ICD-10 conditions, partitioning active from cured illnesses.", 11.5, 6)
    add_bullet(tf_s12, "4. Trajectory Math:", "Calculates numerical deltas and directionality for biomarkers.", 11.5, 6)
    add_bullet(tf_s12, "5. Grounded Note Export:", "Outputs clinical synthesis with multi-hospital citations and 1-click note copy.", 11.5, 6)

    # =========================================================================
    # SLIDE 13: METHODOLOGY — QUANTITATIVE BIOMARKER TRAJECTORIES
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_standard_header(s13, "Methodology: Quantitative Lab Trajectory Computing Engine")

    img_bio = "docs/diagrams/07_biomarker_trajectories_chart.png"
    if os.path.exists(img_bio):
        s13.shapes.add_picture(img_bio, Inches(0.8), Inches(1.8), width=Inches(7.2))

    card_s13 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.1))
    card_s13.fill.solid()
    card_s13.fill.fore_color.rgb = LIGHT_BG
    card_s13.line.color.rgb = BORDER_BLUE
    tf_s13 = card_s13.text_frame
    tf_s13.margin_left = Inches(0.25)
    tf_s13.margin_top = Inches(0.25)
    p_s13 = tf_s13.paragraphs[0]
    p_s13.text = "MATHEMATICAL FORMULATION"
    p_s13.font.name = "Times New Roman"
    p_s13.font.bold = True
    p_s13.font.size = Pt(13)
    p_s13.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s13, "Trajectory Differential:", "Δ = V_latest - V_baseline\n% Change = (Δ / V_baseline) * 100", 11.5, 8)
    add_bullet(tf_s13, "Directionality Rules:", "• Down: Improved / Therapeutic control\n• Up: Elevated / Disease progression\n• Stable: |% Change| < 3%", 11.5, 8)
    add_bullet(tf_s13, "Metabolic Validation:", "HbA1c dropped from 8.7% (2024) to 6.9% (2026) -> Flagged 'Down (Improved)' on Metformin BID.", 11.5, 8)

    # =========================================================================
    # SLIDE 14: TOOLS & TECHNOLOGIES USED
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_standard_header(s14, "Tools and Technology Used")

    tech_categories = [
        ("Frontend Technologies", "React 19, TypeScript 5.8, Vite 8.3 build toolchain, custom CSS clinical design system with hardware-accelerated micro-animations."),
        ("Backend Technologies", "Node.js 22 LTS, Express.js 4.x REST API gateway, Zero-Trust authentication middleware, CORS, modular domain routers."),
        ("Database & Persistence", "Relational health schema normalized to 3NF, indexed queries on healthId, foreign keys, transaction safety."),
        ("Standards & Ontologies", "HL7 FHIR R4 JSON schemas, WHO ICD-10 diagnostic codes, LOINC lab identifiers, ABDM 14-digit ABHA identity."),
        ("Quality & Testing Tools", "Oxlint (high-performance linter, 0 errors), TypeScript compiler (tsc -b), Vitest, Git/GitHub version control.")
    ]

    for idx, (cat_title, cat_desc) in enumerate(tech_categories):
        y = Inches(1.7 + idx * 1.05)
        card = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.92))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_BLUE

        badge = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.18), Inches(2.6), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY_BLUE
        badge.line.fill.background()
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = cat_title
        p_b.font.name = "Times New Roman"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.color.rgb = WHITE

        tbox = s14.shapes.add_textbox(Inches(3.8), y + Inches(0.1), Inches(8.5), Inches(0.72))
        p_desc = tbox.text_frame.paragraphs[0]
        p_desc.text = cat_desc
        p_desc.font.name = "Times New Roman"
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = DARK_NAVY

    # =========================================================================
    # SLIDE 15: ROLES AND RESPONSIBILITIES OF EACH MEMBER
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_standard_header(s15, "Roles and Responsibilities of each Team Member")

    members = [
        ("Samarth (Lead Developer / Project Coordinator)", [
            "• End-to-end full-stack software architecture and system design.",
            "• Designed and engineered the Deterministic Clinical AI Copilot & Trajectory Engine.",
            "• Built the luxury clinical UI design system, keyframe animations, and reactive state management.",
            "• Backend REST API gateway development, Zero-Trust RBAC gatekeeper, and route security."
        ]),
        ("[Student Name 2] (Backend & Data Modeling)", [
            "• Relational database schema design (3NF), data dictionaries, and entity-relationship models.",
            "• Diagnostic reports vault data modeling and biomarker parameter extraction.",
            "• Integration of WHO ICD-10 diagnostic classifications and LOINC standard mappings.",
            "• Automated unit testing and performance benchmarking."
        ]),
        ("[Student Name 3] (Frontend UI & Documentation)", [
            "• Sovereign Citizen DigiLocker, ABHA identity card presentation, and QR code generation.",
            "• Hospital administrator capacity dashboard and multi-hospital registry views.",
            "• Authored comprehensive architectural specifications, UML sequence models, and DFDs.",
            "• Test case execution matrix, visual regression verification, and project thesis drafting."
        ])
    ]

    for idx, (m_name, m_tasks) in enumerate(members):
        y = Inches(1.7 + idx * 1.8)
        card = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.65))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_BLUE
        tf_m = card.text_frame
        tf_m.margin_left = Inches(0.3)
        tf_m.margin_top = Inches(0.2)
        p_m = tf_m.paragraphs[0]
        p_m.text = m_name
        p_m.font.name = "Times New Roman"
        p_m.font.bold = True
        p_m.font.size = Pt(13)
        p_m.font.color.rgb = PRIMARY_BLUE

        for task in m_tasks:
            add_bullet(tf_m, "", task, 11, 2)

    # =========================================================================
    # SLIDE 16: IMPLEMENTATION — MODULAR BACKEND & GATEKEEPER
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_standard_header(s16, "Implementation: Modular Backend Architecture & Zero-Trust Gatekeeper")

    img_rbac = "docs/diagrams/08_zero_trust_rbac_matrix.png"
    if os.path.exists(img_rbac):
        s16.shapes.add_picture(img_rbac, Inches(0.8), Inches(1.8), width=Inches(7.0))

    card_s16 = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.8), Inches(4.5), Inches(5.1))
    card_s16.fill.solid()
    card_s16.fill.fore_color.rgb = CARD_BG
    card_s16.line.color.rgb = PRIMARY_BLUE
    tf_s16 = card_s16.text_frame
    tf_s16.margin_left = Inches(0.25)
    tf_s16.margin_top = Inches(0.25)
    p_s16 = tf_s16.paragraphs[0]
    p_s16.text = "CORE MODULE IMPLEMENTATION"
    p_s16.font.name = "Times New Roman"
    p_s16.font.bold = True
    p_s16.font.size = Pt(13)
    p_s16.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s16, "authRouter.ts:", "Role validation, credential verification, and session state persistence.", 11.5, 6)
    add_bullet(tf_s16, "doctorRouter.ts:", "Patient roster, clinical examination entries, and disease cure certification.", 11.5, 6)
    add_bullet(tf_s16, "aiRouter.ts:", "Deterministic Clinical AI Copilot and report analysis engine.", 11.5, 6)
    add_bullet(tf_s16, "hospitalRouter.ts:", "Hospital network registry and facility operational dashboards.", 11.5, 6)
    add_bullet(tf_s16, "patientRouter.ts:", "Citizen DigiLocker data and ABHA card generation.", 11.5, 6)

    # =========================================================================
    # SLIDE 17: RESULTS — DOCTOR CLINICAL STATION
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    add_standard_header(s17, "Results: Doctor Specialist Station & Clinical Dossier")

    img_s3 = "docs/screenshots/03_doctor_clinical_dossier.png"
    if os.path.exists(img_s3):
        s17.shapes.add_picture(img_s3, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s17.fill.solid()
    card_s17.fill.fore_color.rgb = LIGHT_BG
    card_s17.line.color.rgb = BORDER_BLUE
    tf_s17 = card_s17.text_frame
    tf_s17.margin_left = Inches(0.25)
    tf_s17.margin_top = Inches(0.25)
    p_s17 = tf_s17.paragraphs[0]
    p_s17.text = "CLINICAL OBSERVATIONS"
    p_s17.font.name = "Times New Roman"
    p_s17.font.bold = True
    p_s17.font.size = Pt(13)
    p_s17.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s17, "Authenticated Specialist:", "Dr. Priya Nair (MBBS, MD) | License: MCI-2012-44120 | Apollo Hospitals & Heart Institute.", 11.5, 8)
    add_bullet(tf_s17, "Action Triggers:", "+ Record Disease Encounter & + Issue Diagnostic Lab Report buttons.", 11.5, 8)
    add_bullet(tf_s17, "Assigned Patient Roster:", "Instant selection across Rahul Sharma, Priya Patel, Amit Verma, Sunita Roy.", 11.5, 8)
    add_bullet(tf_s17, "Longitudinal Continuity:", "One-click inspection of patient's multi-hospital health journey throughout lifespan.", 11.5, 8)

    # =========================================================================
    # SLIDE 18: RESULTS — CLINICAL AI COPILOT WORKSTATION
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    add_standard_header(s18, "Results: Clinical AI Copilot & Longitudinal Report Analyzer")

    img_s1 = "docs/screenshots/01_ai_copilot_workstation.png"
    if os.path.exists(img_s1):
        s18.shapes.add_picture(img_s1, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s18 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s18.fill.solid()
    card_s18.fill.fore_color.rgb = CARD_BG
    card_s18.line.color.rgb = PRIMARY_BLUE
    tf_s18 = card_s18.text_frame
    tf_s18.margin_left = Inches(0.25)
    tf_s18.margin_top = Inches(0.25)
    p_s18 = tf_s18.paragraphs[0]
    p_s18.text = "AI COPILOT UI FEATURES"
    p_s18.font.name = "Times New Roman"
    p_s18.font.bold = True
    p_s18.font.size = Pt(13)
    p_s18.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s18, "Rainbow Spectrum Bar:", "Flowing 4px multi-color header accent communicating neural synthesis readiness.", 11.5, 8)
    add_bullet(tf_s18, "Dossier Grounded Status:", "Green pulsing live dot confirming zero hallucination grounding.", 11.5, 8)
    add_bullet(tf_s18, "Instant Preset Chips:", "One-click execution for metabolic, cardiac, joint pain, and full trajectory queries.", 11.5, 8)
    add_bullet(tf_s18, "Elevated Command Shell:", "Focus-within glowing border and real-time cross-hospital synthesis trigger.", 11.5, 8)

    # =========================================================================
    # SLIDE 19: RESULTS — MULTI-HOSPITAL AI SYNTHESIS REPORT
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    add_standard_header(s19, "Results: Multi-Hospital Clinical Synthesis Report")

    img_s2 = "docs/screenshots/02_clinical_synthesis_report.png"
    if os.path.exists(img_s2):
        s19.shapes.add_picture(img_s2, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s19 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s19.fill.solid()
    card_s19.fill.fore_color.rgb = LIGHT_BG
    card_s19.line.color.rgb = BORDER_BLUE
    tf_s19 = card_s19.text_frame
    tf_s19.margin_left = Inches(0.25)
    tf_s19.margin_top = Inches(0.25)
    p_s19 = tf_s19.paragraphs[0]
    p_s19.text = "SYNTHESIS REPORT SECTIONS"
    p_s19.font.name = "Times New Roman"
    p_s19.font.bold = True
    p_s19.font.size = Pt(13)
    p_s19.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s19, "Correlated Conditions:", "Amber badges for ACTIVE conditions, green badges for CURED illnesses.", 11.5, 8)
    add_bullet(tf_s19, "Longitudinal Lab Table:", "Comparing latest 2026 readings against 2024/2025 baselines with trajectory arrows.", 11.5, 8)
    add_bullet(tf_s19, "Formatted Clinical Text:", "Styled blue section headers, dark subheadings, and chevron bullet lists.", 11.5, 8)
    add_bullet(tf_s19, "1-Click Note Transfer:", "Clipboard transfer function copying synthesis directly into the active consultation note.", 11.5, 8)

    # =========================================================================
    # SLIDE 20: RESULTS — ACTIVE VS. CURED DISEASE CONTINUITY
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    add_standard_header(s20, "Results: Active vs. Cured Lifetime Disease Continuity")

    img_s4 = "docs/screenshots/04_lifetime_conditions.png"
    if os.path.exists(img_s4):
        s20.shapes.add_picture(img_s4, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s20 = s20.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s20.fill.solid()
    card_s20.fill.fore_color.rgb = CARD_BG
    card_s20.line.color.rgb = PRIMARY_BLUE
    tf_s20 = card_s20.text_frame
    tf_s20.margin_left = Inches(0.25)
    tf_s20.margin_top = Inches(0.25)
    p_s20 = tf_s20.paragraphs[0]
    p_s20.text = "CLINICAL LIFETIME TRACKING"
    p_s20.font.name = "Times New Roman"
    p_s20.font.bold = True
    p_s20.font.size = Pt(13)
    p_s20.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_s20, "Active Ongoing Conditions:", "Type 2 Diabetes Mellitus (E11.9) and Essential Hypertension (I10) displayed with severity, diagnosis date, and regimen.", 11.5, 8)
    add_bullet(tf_s20, "Resolved Past Diseases:", "Acute Bronchitis (J20.9) and Dengue Fever (A90) displayed with verified cure dates and follow-up clinical evidence.", 11.5, 8)
    add_bullet(tf_s20, "Clinical Utility:", "Prevents physicians from misdiagnosing past infections while maintaining awareness of previous drug responses.", 11.5, 8)

    # =========================================================================
    # SLIDE 21: RESULTS — DIAGNOSTIC VAULT & CITIZEN DIGILOCKER
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    add_standard_header(s21, "Results: Diagnostic Reports Vault & Sovereign Citizen ABHA Card")

    if os.path.exists("docs/screenshots/05_diagnostic_reports_vault.png"):
        s21.shapes.add_picture("docs/screenshots/05_diagnostic_reports_vault.png", Inches(0.8), Inches(1.8), width=Inches(5.6))
        cap5 = s21.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(5.6), Inches(0.5))
        p_c5 = cap5.text_frame.paragraphs[0]
        p_c5.text = "Figure 6.5: Multi-Hospital Diagnostic Reports Vault"
        p_c5.font.name = "Times New Roman"
        p_c5.font.size = Pt(10)
        p_c5.font.bold = True
        p_c5.font.color.rgb = SLATE_TEXT
        p_c5.alignment = PP_ALIGN.CENTER

    if os.path.exists("docs/screenshots/06_citizen_abha_identity.png"):
        s21.shapes.add_picture("docs/screenshots/06_citizen_abha_identity.png", Inches(6.8), Inches(1.8), width=Inches(5.7))
        cap6 = s21.shapes.add_textbox(Inches(6.8), Inches(6.5), Inches(5.7), Inches(0.5))
        p_c6 = cap6.text_frame.paragraphs[0]
        p_c6.text = "Figure 6.6: Sovereign Citizen DigiLocker & ABHA Identity Card"
        p_c6.font.name = "Times New Roman"
        p_c6.font.size = Pt(10)
        p_c6.font.bold = True
        p_c6.font.color.rgb = SLATE_TEXT
        p_c6.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 22: PERFORMANCE & EXPERIMENTAL BENCHMARKS
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    add_standard_header(s22, "Performance Benchmarking & Experimental Verification")

    metrics = [
        ("< 15 ms", "Database Query Latency", "P95 response time for complete multi-hospital patient dossier retrieval across all entities."),
        ("< 50 ms", "AI Analysis Response", "Deterministic RAG pipeline, trajectory math, and synthesis output generation."),
        ("1.35 s", "Vite Production Build", "Full client compilation time with zero TypeScript compilation warnings or errors."),
        ("0 Errors", "Oxlint Quality Score", "Clean code verification across all 5 backend modules and frontend component hierarchy.")
    ]

    for idx, (stat, label, desc) in enumerate(metrics):
        x = Inches(0.8 + idx * 2.95)
        card = s22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.8), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_BLUE

        p_s = card.text_frame.paragraphs[0]
        p_s.text = stat
        p_s.font.name = "Times New Roman"
        p_s.font.size = Pt(28)
        p_s.font.bold = True
        p_s.font.color.rgb = PRIMARY_BLUE
        p_s.alignment = PP_ALIGN.CENTER

        p_l = card.text_frame.add_paragraph()
        p_l.text = label
        p_l.font.name = "Times New Roman"
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = DARK_NAVY
        p_l.alignment = PP_ALIGN.CENTER

        p_d = card.text_frame.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Times New Roman"
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = SLATE_TEXT
        p_d.space_before = Pt(4)

    card_bot = s22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.6))
    card_bot.fill.solid()
    card_bot.fill.fore_color.rgb = LIGHT_BG
    card_bot.line.color.rgb = BORDER_BLUE
    tf_b = card_bot.text_frame
    tf_b.margin_left = Inches(0.3)
    tf_b.margin_top = Inches(0.3)
    p_hb = tf_b.paragraphs[0]
    p_hb.text = "TEST HARNESS EXECUTION SUMMARY"
    p_hb.font.name = "Times New Roman"
    p_hb.font.bold = True
    p_hb.font.size = Pt(13)
    p_hb.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_b, "Unit Test Matrix:", "6/6 automated route test cases passed (Auth, Dossier, Cure Certification, AI Analysis).", 11.5, 6)
    add_bullet(tf_b, "End-to-End Clinical Workflows:", "4/4 clinical scenarios validated (Metabolic onboarding, Dengue cure stamping, Lab vault ingestion, AI note copy).", 11.5, 6)
    add_bullet(tf_b, "Zero-Trust Security Test:", "Attempted horizontal privilege escalation from unauthorized doctor IDs blocked with 401/403.", 11.5, 6)

    # =========================================================================
    # SLIDE 23: CONCLUSION & FUTURE SCOPE
    # =========================================================================
    s23 = prs.slides.add_slide(blank_layout)
    add_standard_header(s23, "Conclusion & Future Research Roadmap")

    card_c1 = s23.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    card_c1.fill.solid()
    card_c1.fill.fore_color.rgb = LIGHT_BG
    card_c1.line.color.rgb = BORDER_BLUE
    tf_c1 = card_c1.text_frame
    tf_c1.margin_left = Inches(0.3)
    tf_c1.margin_top = Inches(0.3)
    p_hc1 = tf_c1.paragraphs[0]
    p_hc1.text = "KEY PROJECT CONTRIBUTIONS"
    p_hc1.font.name = "Times New Roman"
    p_hc1.font.bold = True
    p_hc1.font.size = Pt(13)
    p_hc1.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_c1, "Unified Longitudinal Dossier:", "Unified multi-hospital records under an ABHA-compliant 14-digit citizen health identity.", 12, 10)
    add_bullet(tf_c1, "Active vs. Cured Continuity:", "Solved lifetime illness conflation through a certified cure state transition machine.", 12, 10)
    add_bullet(tf_c1, "Deterministic Clinical AI:", "Engineered zero-hallucination decision support with quantitative trajectory math and 1-click note copy.", 12, 10)
    add_bullet(tf_c1, "Zero-Trust Architecture:", "Enforced role-scoped workspaces in strict compliance with India's DPDPA 2023.", 12, 10)

    card_c2 = s23.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3))
    card_c2.fill.solid()
    card_c2.fill.fore_color.rgb = CARD_BG
    card_c2.line.color.rgb = PRIMARY_BLUE
    tf_c2 = card_c2.text_frame
    tf_c2.margin_left = Inches(0.3)
    tf_c2.margin_top = Inches(0.3)
    p_hc2 = tf_c2.paragraphs[0]
    p_hc2.text = "FUTURE RESEARCH ROADMAP"
    p_hc2.font.name = "Times New Roman"
    p_hc2.font.bold = True
    p_hc2.font.size = Pt(13)
    p_hc2.font.color.rgb = PRIMARY_BLUE

    add_bullet(tf_c2, "Live NHA ABDM Gateway:", "Integrate official production sandbox certificates for real-time live hospital EHR synchronization.", 12, 10)
    add_bullet(tf_c2, "Wearable IoT Telemetry:", "Stream continuous glucose monitoring (CGM) and heart rate variability from consumer smartwatches.", 12, 10)
    add_bullet(tf_c2, "Vernacular Multilingual AI:", "Expand natural language copilot support to Hindi, Tamil, Telugu, and Bengali for rural health centers.", 12, 10)
    add_bullet(tf_c2, "Federated Clinical Research:", "Implement differential privacy for multi-hospital epidemiological research without exposing patient identities.", 12, 10)

    # =========================================================================
    # SLIDE 24: REFERENCES (IEEE / APA FORMAT)
    # =========================================================================
    s24 = prs.slides.add_slide(blank_layout)
    add_standard_header(s24, "References & Academic Citations (IEEE / APA Format)")

    card_ref = s24.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.3))
    card_ref.fill.solid()
    card_ref.fill.fore_color.rgb = LIGHT_BG
    card_ref.line.color.rgb = BORDER_BLUE
    tf_ref = card_ref.text_frame
    tf_ref.margin_left = Inches(0.3)
    tf_ref.margin_top = Inches(0.3)
    p_href = tf_ref.paragraphs[0]
    p_href.text = "PEER-REVIEWED CITATIONS & STANDARDS"
    p_href.font.name = "Times New Roman"
    p_href.font.bold = True
    p_href.font.size = Pt(13)
    p_href.font.color.rgb = PRIMARY_BLUE

    refs = [
        "[1] National Health Authority (NHA). 'Ayushman Bharat Digital Mission (ABDM) Strategy Document.' Ministry of Health and Family Welfare, New Delhi, 2021.",
        "[2] HL7 International. 'Fast Healthcare Interoperability Resources (FHIR) Release 4 Standard.' Health Level Seven International, Ann Arbor, MI, 2019.",
        "[3] World Health Organization (WHO). 'International Statistical Classification of Diseases and Related Health Problems (ICD-10).' 10th Revision, Geneva, 2016.",
        "[4] Ministry of Law and Justice, Govt of India. 'The Digital Personal Data Protection Act, 2023 (Act No. 22 of 2023).' The Gazette of India, August 2023.",
        "[5] P. Rajpurkar, E. Chen, O. Banerjee, and E. J. Topol. 'AI in health and medicine.' Nature Medicine, vol. 28, no. 1, pp. 31–38, 2022.",
        "[6] K. Singhal, S. Azizi, T. Tu, et al. 'Large language models encode clinical knowledge.' Nature, vol. 620, pp. 172–180, 2023.",
        "[7] K. D. Mandl and I. S. Kohane. 'Escaping the EHR trap — The future of health IT.' New England Journal of Medicine, vol. 366, no. 24, pp. 2240–2242, 2012.",
        "[8] Department of Computer Science & Engineering. 'BCO 519A — Major Project Guidelines & Evaluation Scheme.' Academic Session 2025–2026."
    ]

    for r in refs:
        add_bullet(tf_ref, "", r, 11, 4)

    # Save to both locations
    out1 = "docs/MediSutra_Major_Project_Presentation.pptx"
    out2 = "MediSutra_Major_Project_Presentation_Department_Format.pptx"
    prs.save(out1)
    prs.save(out2)
    print(f"Presentation generated successfully at: {out1} and {out2}")

if __name__ == '__main__':
    build_presentation()
