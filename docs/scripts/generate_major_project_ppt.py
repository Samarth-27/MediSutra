import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY_BLUE = RGBColor(2, 132, 199)    # #0284C7
    DARK_NAVY    = RGBColor(15, 23, 42)     # #0F172A
    TEAL_ACCENT  = RGBColor(13, 148, 136)   # #0D9488
    SLATE_TEXT   = RGBColor(51, 65, 85)     # #334155
    LIGHT_BG     = RGBColor(248, 250, 252)  # #F8FAFC
    CARD_BG      = RGBColor(240, 249, 255)  # #F0F9FF
    WHITE        = RGBColor(255, 255, 255)
    BORDER_BLUE  = RGBColor(186, 230, 253)  # #BAE6FD
    GREEN_ACCENT = RGBColor(16, 185, 129)   # #10B981

    def add_header(slide, title_text, category_text="BCO 519A — MAJOR PROJECT DEFENSE"):
        # Top banner accent line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        accent.fill.solid()
        accent.fill.fore_color.rgb = PRIMARY_BLUE
        accent.line.color.rgb = PRIMARY_BLUE

        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.4))
        p_cat = cat_box.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_BLUE

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.8))
        p_t = t_box.text_frame.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_NAVY

        # Subtle bottom line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_BLUE
        line.line.color.rgb = BORDER_BLUE

    def add_bullet_point(tf, bold_txt, body_txt, pt_size=13, space_after=8):
        p = tf.add_paragraph()
        p.space_after = Pt(space_after)
        if bold_txt:
            r1 = p.add_run()
            r1.text = bold_txt + " "
            r1.font.bold = True
            r1.font.size = Pt(pt_size)
            r1.font.color.rgb = DARK_NAVY
        r2 = p.add_run()
        r2.text = body_txt
        r2.font.size = Pt(pt_size)
        r2.font.color.rgb = SLATE_TEXT

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_NAVY
    bg1.line.fill.background()

    # Top accent line
    top1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    top1.fill.solid()
    top1.fill.fore_color.rgb = PRIMARY_BLUE
    top1.line.fill.background()

    # Tag Badge
    tag1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.9), Inches(4.5), Inches(0.45))
    tag1.fill.solid()
    tag1.fill.fore_color.rgb = PRIMARY_BLUE
    tag1.line.fill.background()
    p_tag1 = tag1.text_frame.paragraphs[0]
    p_tag1.text = "BCO 519A — MAJOR PROJECT DISSERTATION DEFENSE"
    p_tag1.font.size = Pt(10)
    p_tag1.font.bold = True
    p_tag1.alignment = PP_ALIGN.CENTER
    p_tag1.font.color.rgb = WHITE

    # Project Title
    t_box1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(1.6))
    p_title1 = t_box1.text_frame.paragraphs[0]
    p_title1.text = "MEDISUTRA"
    p_title1.font.size = Pt(44)
    p_title1.font.bold = True
    p_title1.font.color.rgb = WHITE

    p_sub1 = t_box1.text_frame.add_paragraph()
    p_sub1.text = "A Unified Longitudinal Personal Health Intelligence & Federated Clinical Copilot Platform"
    p_sub1.font.size = Pt(20)
    p_sub1.font.color.rgb = RGBColor(186, 230, 253)

    # Sub-tagline
    p_tagline = t_box1.text_frame.add_paragraph()
    p_tagline.text = "Tagline: One Person. One Health Identity. One Complete Health Journey."
    p_tagline.font.size = Pt(13)
    p_tagline.font.color.rgb = RGBColor(148, 163, 184)
    p_tagline.space_before = Pt(8)

    # Bottom Metadata Card
    card_meta = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.5), Inches(10.9), Inches(2.2))
    card_meta.fill.solid()
    card_meta.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card_meta.line.color.rgb = RGBColor(51, 65, 85)

    tf_meta = card_meta.text_frame
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "PROJECT METADATA & COMMITTEE SUBMISSION"
    p_m1.font.size = Pt(11)
    p_m1.font.bold = True
    p_m1.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_meta, "Submitted By:", "Samarth (Lead Developer / Project Coordinator) | Enrollment: [Roll No]", 12, 4)
    add_bullet_point(tf_meta, "Team Members:", "[Student Name 2], [Student Name 3] | B.Tech Computer Science & Engineering", 12, 4)
    add_bullet_point(tf_meta, "Under Supervision of:", "[Faculty Guide Name & Designation] | Dept of Computer Science & Engineering", 12, 4)
    add_bullet_point(tf_meta, "Department & Institution:", "Department of Computer Science & Engineering, [Institution / University Name]", 12, 4)

    # =========================================================================
    # SLIDE 2: PRESENTATION OUTLINE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Executive Presentation Outline & Flow")

    agenda_items = [
        ("01", "Introduction & Clinical Problem", "Health data fragmentation dilemma in India and motivation"),
        ("02", "Project Objectives & Scope", "7 core deliverables aligned with Course Outcomes CO1–CO5"),
        ("03", "Literature Review & Research Gap", "Comparative analysis with Epic, Apple Health, ABDM PHR apps"),
        ("04", "System Architecture & DFDs", "4-tier architecture, DFD Level 0/1/2, HRC network topology"),
        ("05", "Database Design & ER Modeling", "Relational schema 3NF, data dictionaries, state machines"),
        ("06", "Deterministic Clinical AI Engine", "Cross-hospital RAG, biomarker trajectory math, zero hallucination"),
        ("07", "Zero-Trust RBAC & Security", "Role gatekeeper, citizen DigiLocker, DPDPA 2023 compliance"),
        ("08", "Live Results & System Screenshots", "Doctor workstation, AI report analyzer, active vs cured diseases"),
        ("09", "Verification & Performance", "Automated test matrix, latency benchmarks, case studies"),
        ("10", "Conclusion & Future Roadmap", "Course Outcome attainment, limitations, and future research")
    ]

    for idx, (num, title, desc) in enumerate(agenda_items):
        col = idx // 5
        row = idx % 5
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.7 + row * 1.05)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_BLUE

        # Number circle
        num_circle = s2.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.15), Inches(0.6), Inches(0.6))
        num_circle.fill.solid()
        num_circle.fill.fore_color.rgb = PRIMARY_BLUE
        num_circle.line.fill.background()
        p_n = num_circle.text_frame.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(11)
        p_n.font.bold = True
        p_n.font.color.rgb = WHITE
        p_n.alignment = PP_ALIGN.CENTER

        # Text
        txt_box = s2.shapes.add_textbox(x + Inches(0.85), y + Inches(0.08), Inches(4.6), Inches(0.75))
        p_t = txt_box.text_frame.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_NAVY

        p_d = txt_box.text_frame.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = SLATE_TEXT

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT & MOTIVATION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Clinical Problem Statement: Healthcare Data Fragmentation")

    # Left Card: The Real-World Dilemma
    card_p1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    card_p1.fill.solid()
    card_p1.fill.fore_color.rgb = LIGHT_BG
    card_p1.line.color.rgb = BORDER_BLUE
    tf_p1 = card_p1.text_frame
    tf_p1.margin_left = Inches(0.3)
    tf_p1.margin_top = Inches(0.3)
    p_h1 = tf_p1.paragraphs[0]
    p_h1.text = "THE FRAGMENTATION DILEMMA IN INDIA"
    p_h1.font.bold = True
    p_h1.font.size = Pt(13)
    p_h1.font.color.rgb = RGBColor(220, 38, 38)

    add_bullet_point(tf_p1, "Paper Record Burden:", "Patients carry physical paper files containing ultrasound films, ECGs, prescriptions, and lab tests across hospital visits.", 12, 10)
    add_bullet_point(tf_p1, "Redundant Investigations:", "Due to inaccessible past records, doctors re-order expensive lab tests (HbA1c, CT scans), imposing unnecessary costs and wasting laboratory capacity.", 12, 10)
    add_bullet_point(tf_p1, "Clinical Blindspots:", "Doctors have no immediate visibility into historical adverse drug reactions, past hospitalizations, or organ panel trajectories.", 12, 10)
    add_bullet_point(tf_p1, "Conflation of Acute & Chronic:", "Cured acute illnesses (Dengue, Bronchitis) are conflated with ongoing chronic diseases (Diabetes, Hypertension), confusing clinical context.", 12, 10)

    # Right Card: The Solution Needed
    card_p2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3))
    card_p2.fill.solid()
    card_p2.fill.fore_color.rgb = CARD_BG
    card_p2.line.color.rgb = PRIMARY_BLUE
    tf_p2 = card_p2.text_frame
    tf_p2.margin_left = Inches(0.3)
    tf_p2.margin_top = Inches(0.3)
    p_h2 = tf_p2.paragraphs[0]
    p_h2.text = "MEDISUTRA RESEARCH & ENGINEERING RESPONSE"
    p_h2.font.bold = True
    p_h2.font.size = Pt(13)
    p_h2.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_p2, "Unified Health Identity:", "Standardized 14-digit ABHA/ABDM citizen identifier linking multi-hospital records to one sovereign profile.", 12, 10)
    add_bullet_point(tf_p2, "Federated Hospital Network:", "Connecting Apollo, Fortis, Max, AIIMS, Dr. Lal PathLabs, and Metropolis into an interoperable clinical graph.", 12, 10)
    add_bullet_point(tf_p2, "Certified Disease Lifecycle:", "Explicit state machine separating active ongoing chronic conditions from certified cured historical illnesses.", 12, 10)
    add_bullet_point(tf_p2, "Deterministic Clinical AI:", "Decision-support copilot that correlates symptoms with patient records and computes lab trajectories with zero hallucinations.", 12, 10)
    add_bullet_point(tf_p2, "Zero-Trust Security:", "Role-scoped workspaces compliant with India's Digital Personal Data Protection Act (DPDPA 2023).", 12, 10)

    # =========================================================================
    # SLIDE 4: MAJOR PROJECT OBJECTIVES & COURSE OUTCOMES
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Project Objectives & BCO 519A Course Outcome Attainment")

    objs = [
        ("CO1: Problem Formulation", "Formulate an engineering computing solution for national health data fragmentation using ABDM standards and condition-centric longitudinal data models."),
        ("CO2: System Design & Tooling", "Architect a 4-tier modular web platform using React 19, TypeScript, Node.js/Express, and normalized relational database schemas with ER modeling."),
        ("CO3: Implementation & Testing", "Implement clean modular backend micro-services and reactive frontend workstations with 0 compiler/lint errors and rigorous E2E test verification."),
        ("CO4: Ethics & Project Mgmt", "Enforce Zero-Trust security and data sovereignty adhering to India's DPDPA 2023 and agile GitHub version control discipline (Samarth-27/MediSutra)."),
        ("CO5: Technical Defense", "Deliver publication-grade architectural documentation, interactive live system demonstrations, visual recordings, and this dissertation defense.")
    ]

    for idx, (co_title, co_desc) in enumerate(objs):
        y = Inches(1.7 + idx * 1.05)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.92))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_BLUE

        badge = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y + Inches(0.18), Inches(2.6), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY_BLUE
        badge.line.fill.background()
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = co_title
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.color.rgb = WHITE

        tbox = s4.shapes.add_textbox(Inches(3.8), y + Inches(0.1), Inches(8.5), Inches(0.72))
        p_desc = tbox.text_frame.paragraphs[0]
        p_desc.text = co_desc
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = DARK_NAVY

    # =========================================================================
    # SLIDE 5: SYSTEM ARCHITECTURE (4-TIER)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Enterprise 4-Tier System Architecture")

    # Image Left
    img_arch = "docs/diagrams/01_system_architecture_4_tier.png"
    if os.path.exists(img_arch):
        s5.shapes.add_picture(img_arch, Inches(0.8), Inches(1.7), width=Inches(6.8))

    # Text Right
    card_a = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_a.fill.solid()
    card_a.fill.fore_color.rgb = LIGHT_BG
    card_a.line.color.rgb = BORDER_BLUE
    tf_a = card_a.text_frame
    tf_a.margin_left = Inches(0.25)
    tf_a.margin_top = Inches(0.25)
    p_ha = tf_a.paragraphs[0]
    p_ha.text = "ARCHITECTURAL HIGHLIGHTS"
    p_ha.font.bold = True
    p_ha.font.size = Pt(13)
    p_ha.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_a, "Tier 1: Presentation Layer", "React 19 + TypeScript + Vite. Role-partitioned clinical workstations for Doctors, Hospital Admins, and Citizens.", 11.5, 8)
    add_bullet_point(tf_a, "Tier 2: API Gateway Layer", "Express.js REST gateway with Zero-Trust RBAC middleware, session persistence, and standardized JSON envelopes.", 11.5, 8)
    add_bullet_point(tf_a, "Tier 3: Clinical Intelligence", "Deterministic RAG pipeline, quantitative trajectory computing engine, and disease lifecycle state machine.", 11.5, 8)
    add_bullet_point(tf_a, "Tier 4: Persistence Layer", "Relational health models, 14-digit ABHA patient registry, and immutable audit logs.", 11.5, 8)

    # =========================================================================
    # SLIDE 6: DATA FLOW ARCHITECTURE (DFD LEVEL 0 & LEVEL 1)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Data Flow Architecture: Context Diagram & Core Subsystems")

    img_dfd0 = "docs/diagrams/02_dfd_level_0_context.png"
    img_dfd1 = "docs/diagrams/03_dfd_level_1_subsystems.png"

    if os.path.exists(img_dfd0):
        s6.shapes.add_picture(img_dfd0, Inches(0.8), Inches(1.8), width=Inches(5.6))
        cap0 = s6.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(5.6), Inches(0.5))
        p_c0 = cap0.text_frame.paragraphs[0]
        p_c0.text = "Figure 4.2: DFD Level 0 Context Level Diagram"
        p_c0.font.size = Pt(10)
        p_c0.font.bold = True
        p_c0.font.color.rgb = SLATE_TEXT
        p_c0.alignment = PP_ALIGN.CENTER

    if os.path.exists(img_dfd1):
        s6.shapes.add_picture(img_dfd1, Inches(6.8), Inches(1.8), width=Inches(5.7))
        cap1 = s6.shapes.add_textbox(Inches(6.8), Inches(6.5), Inches(5.7), Inches(0.5))
        p_c1 = cap1.text_frame.paragraphs[0]
        p_c1.text = "Figure 4.3: DFD Level 1 Core Subsystems Decomposition"
        p_c1.font.size = Pt(10)
        p_c1.font.bold = True
        p_c1.font.color.rgb = SLATE_TEXT
        p_c1.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 7: FEDERATED HOSPITAL NETWORK (HRC TOPOLOGY)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Federated Human Report Center (HRC) Network Topology")

    img_hrc = "docs/diagrams/04_hrc_network_topology.png"
    if os.path.exists(img_hrc):
        s7.shapes.add_picture(img_hrc, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_hrc = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_hrc.fill.solid()
    card_hrc.fill.fore_color.rgb = CARD_BG
    card_hrc.line.color.rgb = PRIMARY_BLUE
    tf_hrc = card_hrc.text_frame
    tf_hrc.margin_left = Inches(0.25)
    tf_hrc.margin_top = Inches(0.25)
    p_hh = tf_hrc.paragraphs[0]
    p_hh.text = "NETWORK FEDERATION FEATURES"
    p_hh.font.bold = True
    p_hh.font.size = Pt(13)
    p_hh.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_hrc, "Central Registry Hub:", "Maintains unified patient identities anchored to the 14-digit ABHA number.", 11.5, 8)
    add_bullet_point(tf_hrc, "Hospital Nodes:", "Apollo Hospitals, Fortis Memorial, Max Super Speciality, AIIMS New Delhi operate authenticated clinical stations.", 11.5, 8)
    add_bullet_point(tf_hrc, "Diagnostic Nodes:", "Dr. Lal PathLabs and Metropolis Healthcare ingest laboratory test panels directly into the central vault.", 11.5, 8)
    add_bullet_point(tf_hrc, "Decentralized Continuity:", "Hospitals retain institutional autonomy while contributing standardized encounter data to the patient's sovereign health graph.", 11.5, 8)

    # =========================================================================
    # SLIDE 8: DATABASE DESIGN & ER MODEL
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Database Architecture & Entity-Relationship (ER) Model")

    img_er = "docs/diagrams/10_conceptual_er_diagram.png"
    if os.path.exists(img_er):
        s8.shapes.add_picture(img_er, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_er = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_er.fill.solid()
    card_er.fill.fore_color.rgb = LIGHT_BG
    card_er.line.color.rgb = BORDER_BLUE
    tf_er = card_er.text_frame
    tf_er.margin_left = Inches(0.25)
    tf_er.margin_top = Inches(0.25)
    p_her = tf_er.paragraphs[0]
    p_her.text = "3NF RELATIONAL ENTITIES"
    p_her.font.bold = True
    p_her.font.size = Pt(13)
    p_her.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_er, "Patients Table:", "Stores healthId (PK), abhaNumber (14-digit unique), demographic data, and primary facility.", 11.5, 8)
    add_bullet_point(tf_er, "Encounters Table:", "Captures doctorId, hospitalId, ICD-10 diagnosis, clinical severity, and examination notes.", 11.5, 8)
    add_bullet_point(tf_er, "Conditions Table:", "Maintains lifetime disease state (ACTIVE vs. RESOLVED), diagnosedDate, curedDate, and curedEvidence.", 11.5, 8)
    add_bullet_point(tf_er, "Reports & Biomarkers:", "Stores multi-year lab panels and quantitative analyte readings (HbA1c, Glucose, Vitamin D).", 11.5, 8)

    # =========================================================================
    # SLIDE 9: LIFETIME DISEASE LIFECYCLE (ACTIVE VS. CURED)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Clinical Disease Lifecycle State Transition Machine")

    img_sm = "docs/diagrams/05_disease_lifecycle_state_machine.png"
    if os.path.exists(img_sm):
        s9.shapes.add_picture(img_sm, Inches(0.8), Inches(1.8), width=Inches(7.2))

    card_sm = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.1))
    card_sm.fill.solid()
    card_sm.fill.fore_color.rgb = CARD_BG
    card_sm.line.color.rgb = PRIMARY_BLUE
    tf_sm = card_sm.text_frame
    tf_sm.margin_left = Inches(0.25)
    tf_sm.margin_top = Inches(0.25)
    p_hsm = tf_sm.paragraphs[0]
    p_hsm.text = "EXPLICIT CLINICAL SEPARATION"
    p_hsm.font.bold = True
    p_hsm.font.size = Pt(13)
    p_hsm.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_sm, "Active Chronic Conditions:", "Diseases requiring ongoing therapy (Type 2 Diabetes, Hypertension) remain tagged ACTIVE.", 11.5, 8)
    add_bullet_point(tf_sm, "Cure Certification:", "Doctors certify illness resolution by submitting follow-up clinical examination notes and diagnostic evidence.", 11.5, 8)
    add_bullet_point(tf_sm, "Longitudinal History:", "Cured illnesses (Dengue, Bronchitis) transition to RESOLVED and are archived into historical records with verified evidence.", 11.5, 8)

    # =========================================================================
    # SLIDE 10: DETERMINISTIC CLINICAL AI COPILOT
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Deterministic Clinical AI Decision-Support Copilot")

    img_ai = "docs/diagrams/06_clinical_ai_rag_flowchart.png"
    if os.path.exists(img_ai):
        s10.shapes.add_picture(img_ai, Inches(0.8), Inches(1.7), width=Inches(6.8))

    card_ai = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.7), Inches(4.7), Inches(5.3))
    card_ai.fill.solid()
    card_ai.fill.fore_color.rgb = LIGHT_BG
    card_ai.line.color.rgb = BORDER_BLUE
    tf_ai = card_ai.text_frame
    tf_ai.margin_left = Inches(0.25)
    tf_ai.margin_top = Inches(0.25)
    p_hai = tf_ai.paragraphs[0]
    p_hai.text = "ZERO-HALLUCINATION RAG"
    p_hai.font.bold = True
    p_hai.font.size = Pt(13)
    p_hai.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_ai, "1. Query Ingestion:", "Doctors write freeform notes or select instant preset chips (e.g., 'High Fasting Sugar & Tingling').", 11.5, 6)
    add_bullet_point(tf_ai, "2. Dossier Retrieval:", "Fetches multi-hospital encounters, active/cured diseases, and laboratory panels.", 11.5, 6)
    add_bullet_point(tf_ai, "3. Temporal Correlation:", "Maps clinical tokens to ICD-10 conditions across historical calendar years.", 11.5, 6)
    add_bullet_point(tf_ai, "4. Trajectory Computing:", "Mathematically computes biomarker differentials (HbA1c, Glucose, Vitamin D).", 11.5, 6)
    add_bullet_point(tf_ai, "5. Grounded Note Export:", "Outputs structured synthesis with multi-hospital citations and 1-click consultation note insertion.", 11.5, 6)

    # =========================================================================
    # SLIDE 11: QUANTITATIVE BIOMARKER TRAJECTORIES
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Quantitative Laboratory Trajectory Computing Engine")

    img_bio = "docs/diagrams/07_biomarker_trajectories_chart.png"
    if os.path.exists(img_bio):
        s11.shapes.add_picture(img_bio, Inches(0.8), Inches(1.8), width=Inches(7.2))

    card_bio = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.1))
    card_bio.fill.solid()
    card_bio.fill.fore_color.rgb = CARD_BG
    card_bio.line.color.rgb = PRIMARY_BLUE
    tf_bio = card_bio.text_frame
    tf_bio.margin_left = Inches(0.25)
    tf_bio.margin_top = Inches(0.25)
    p_hbio = tf_bio.paragraphs[0]
    p_hbio.text = "MATHEMATICAL FORMULATION"
    p_hbio.font.bold = True
    p_hbio.font.size = Pt(13)
    p_hbio.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_bio, "Trajectory Differential:", "Δ = V_latest - V_baseline\n% Change = (Δ / V_baseline) * 100", 11.5, 8)
    add_bullet_point(tf_bio, "Directionality Rules:", "• Down: Improved / Therapeutic control\n• Up: Elevated / Disease progression\n• Stable: |% Change| < 3%", 11.5, 8)
    add_bullet_point(tf_bio, "Metabolic Validation:", "HbA1c dropped from 8.7% (2024) to 6.9% (2026) -> Flagged 'Down (Improved)' on Metformin BID.", 11.5, 8)

    # =========================================================================
    # SLIDE 12: ZERO-TRUST RBAC & SECURITY ARCHITECTURE
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Zero-Trust Security & Role-Based Access Control (RBAC)")

    img_rbac = "docs/diagrams/08_zero_trust_rbac_matrix.png"
    if os.path.exists(img_rbac):
        s12.shapes.add_picture(img_rbac, Inches(0.8), Inches(1.8), width=Inches(7.0))

    card_sec = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.8), Inches(4.5), Inches(5.1))
    card_sec.fill.solid()
    card_sec.fill.fore_color.rgb = LIGHT_BG
    card_sec.line.color.rgb = BORDER_BLUE
    tf_sec = card_sec.text_frame
    tf_sec.margin_left = Inches(0.25)
    tf_sec.margin_top = Inches(0.25)
    p_hsec = tf_sec.paragraphs[0]
    p_hsec.text = "SECURITY & DPDPA GOVERNANCE"
    p_hsec.font.bold = True
    p_hsec.font.size = Pt(13)
    p_hsec.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_sec, "Doctor Workspace:", "Restricted to assigned patient dossiers. Can record encounters, stamp cures, and query AI copilot.", 11.5, 8)
    add_bullet_point(tf_sec, "Hospital Admin Workspace:", "Restricted to facility capacity and doctor directories. Cannot inspect patient charts.", 11.5, 8)
    add_bullet_point(tf_sec, "Sovereign Citizen DigiLocker:", "Patient maintains full cryptographic ownership of personal and family health records.", 11.5, 8)
    add_bullet_point(tf_sec, "DPDPA 2023 Compliance:", "Purpose limitation, data minimization, and immutable audit trails for every query.", 11.5, 8)

    # =========================================================================
    # SLIDE 13: UML SEQUENCE DIAGRAM
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Clinical Consultation & AI Synthesis Workflow (UML Sequence)")

    img_seq = "docs/diagrams/09_uml_consultation_sequence.png"
    if os.path.exists(img_seq):
        s13.shapes.add_picture(img_seq, Inches(0.8), Inches(1.7), width=Inches(7.0))

    card_seq = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.7), Inches(4.5), Inches(5.3))
    card_seq.fill.solid()
    card_seq.fill.fore_color.rgb = CARD_BG
    card_seq.line.color.rgb = PRIMARY_BLUE
    tf_seq = card_seq.text_frame
    tf_seq.margin_left = Inches(0.25)
    tf_seq.margin_top = Inches(0.25)
    p_hseq = tf_seq.paragraphs[0]
    p_hseq.text = "INTERACTION SEQUENCE"
    p_hseq.font.bold = True
    p_hseq.font.size = Pt(13)
    p_hseq.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_seq, "1. Patient Selection:", "Doctor selects patient from assigned roster.", 11.5, 6)
    add_bullet_point(tf_seq, "2. Dossier Ingestion:", "Gateway queries multi-hospital data store and streams longitudinal record.", 11.5, 6)
    add_bullet_point(tf_seq, "3. Copilot Query:", "Doctor clicks preset chip or types clinical symptoms.", 11.5, 6)
    add_bullet_point(tf_seq, "4. Real-Time Synthesis:", "AI engine correlates symptoms, computes trajectories, and returns report.", 11.5, 6)
    add_bullet_point(tf_seq, "5. Consultation Note Export:", "1-click copy directly inserts structured intelligence into the active note.", 11.5, 6)

    # =========================================================================
    # SLIDE 14: LIVE SCREENSHOT: DOCTOR CLINICAL STATION
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Live System Demonstration: Doctor Specialist Clinical Station")

    img_s3 = "docs/screenshots/03_doctor_clinical_dossier.png"
    if os.path.exists(img_s3):
        s14.shapes.add_picture(img_s3, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s14 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s14.fill.solid()
    card_s14.fill.fore_color.rgb = LIGHT_BG
    card_s14.line.color.rgb = BORDER_BLUE
    tf_s14 = card_s14.text_frame
    tf_s14.margin_left = Inches(0.25)
    tf_s14.margin_top = Inches(0.25)
    p_hs14 = tf_s14.paragraphs[0]
    p_hs14.text = "CLINICAL WORKSTATION CAPABILITIES"
    p_hs14.font.bold = True
    p_hs14.font.size = Pt(13)
    p_hs14.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_s14, "Authenticated Specialist:", "Dr. Priya Nair (MBBS, MD) | License: MCI-2012-44120 | Apollo Hospitals & Heart Institute.", 11.5, 8)
    add_bullet_point(tf_s14, "Action Triggers:", "+ Record Disease Encounter & + Issue Diagnostic Lab Report buttons.", 11.5, 8)
    add_bullet_point(tf_s14, "Assigned Patient Roster:", "Instant selection across Rahul Sharma, Priya Patel, Amit Verma, Sunita Roy.", 11.5, 8)
    add_bullet_point(tf_s14, "Longitudinal Continuity:", "One-click inspection of patient's multi-hospital health journey throughout lifespan.", 11.5, 8)

    # =========================================================================
    # SLIDE 15: LIVE SCREENSHOT: CLINICAL AI COPILOT WORKSTATION
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Live System Demonstration: Clinical AI Copilot & Presets")

    img_s1 = "docs/screenshots/01_ai_copilot_workstation.png"
    if os.path.exists(img_s1):
        s15.shapes.add_picture(img_s1, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s15 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s15.fill.solid()
    card_s15.fill.fore_color.rgb = CARD_BG
    card_s15.line.color.rgb = PRIMARY_BLUE
    tf_s15 = card_s15.text_frame
    tf_s15.margin_left = Inches(0.25)
    tf_s15.margin_top = Inches(0.25)
    p_hs15 = tf_s15.paragraphs[0]
    p_hs15.text = "AI COPILOT DESIGN INNOVATIONS"
    p_hs15.font.bold = True
    p_hs15.font.size = Pt(13)
    p_hs15.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_s15, "Rainbow Spectrum Bar:", "Flowing 4px multi-color header accent communicating neural synthesis readiness.", 11.5, 8)
    add_bullet_point(tf_s15, "Dossier Grounded Status:", "Green pulsing live dot confirming zero hallucination grounding.", 11.5, 8)
    add_bullet_point(tf_s15, "Instant Preset Chips:", "One-click execution for metabolic, cardiac, joint pain, and full trajectory queries.", 11.5, 8)
    add_bullet_point(tf_s15, "Elevated Command Shell:", "Focus-within glowing border and real-time cross-hospital synthesis trigger.", 11.5, 8)

    # =========================================================================
    # SLIDE 16: LIVE SCREENSHOT: MULTI-HOSPITAL AI SYNTHESIS REPORT
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "Live System Demonstration: Clinical Intelligence Synthesis Report")

    img_s2 = "docs/screenshots/02_clinical_synthesis_report.png"
    if os.path.exists(img_s2):
        s16.shapes.add_picture(img_s2, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s16 = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s16.fill.solid()
    card_s16.fill.fore_color.rgb = LIGHT_BG
    card_s16.line.color.rgb = BORDER_BLUE
    tf_s16 = card_s16.text_frame
    tf_s16.margin_left = Inches(0.25)
    tf_s16.margin_top = Inches(0.25)
    p_hs16 = tf_s16.paragraphs[0]
    p_hs16.text = "SYNTHESIS REPORT SECTIONS"
    p_hs16.font.bold = True
    p_hs16.font.size = Pt(13)
    p_hs16.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_s16, "Correlated Conditions:", "Amber badges for ACTIVE conditions, green badges for CURED illnesses.", 11.5, 8)
    add_bullet_point(tf_s16, "Longitudinal Lab Table:", "Comparing latest 2026 readings against 2024/2025 baselines with trajectory arrows.", 11.5, 8)
    add_bullet_point(tf_s16, "Formatted Clinical Text:", "Styled blue section headers, dark subheadings, and chevron bullet lists.", 11.5, 8)
    add_bullet_point(tf_s16, "1-Click Note Transfer:", "Clipboard transfer function copying synthesis directly into the active consultation note.", 11.5, 8)

    # =========================================================================
    # SLIDE 17: LIVE SCREENSHOT: ACTIVE VS. CURED DISEASE CONTINUITY
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "Live System Demonstration: Active vs. Cured Disease Continuity")

    img_s4 = "docs/screenshots/04_lifetime_conditions.png"
    if os.path.exists(img_s4):
        s17.shapes.add_picture(img_s4, Inches(0.8), Inches(1.7), width=Inches(7.2))

    card_s17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.7), Inches(4.3), Inches(5.3))
    card_s17.fill.solid()
    card_s17.fill.fore_color.rgb = CARD_BG
    card_s17.line.color.rgb = PRIMARY_BLUE
    tf_s17 = card_s17.text_frame
    tf_s17.margin_left = Inches(0.25)
    tf_s17.margin_top = Inches(0.25)
    p_hs17 = tf_s17.paragraphs[0]
    p_hs17.text = "CLINICAL LIFETIME TRACKING"
    p_hs17.font.bold = True
    p_hs17.font.size = Pt(13)
    p_hs17.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_s17, "Active Ongoing Conditions:", "Type 2 Diabetes Mellitus (E11.9) and Essential Hypertension (I10) displayed with severity, diagnosis date, and regimen.", 11.5, 8)
    add_bullet_point(tf_s17, "Resolved Past Diseases:", "Acute Bronchitis (J20.9) and Dengue Fever (A90) displayed with verified cure dates and follow-up clinical evidence.", 11.5, 8)
    add_bullet_point(tf_s17, "Clinical Utility:", "Prevents physicians from misdiagnosing past infections while maintaining awareness of previous drug responses.", 11.5, 8)

    # =========================================================================
    # SLIDE 18: LIVE SCREENSHOT: DIAGNOSTIC VAULT & CITIZEN ABHA IDENTITY
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    add_header(s18, "Live System Demonstration: Diagnostic Reports Vault & ABHA Card")

    img_s5 = "docs/screenshots/05_diagnostic_reports_vault.png"
    img_s6 = "docs/screenshots/06_citizen_abha_identity.png"

    if os.path.exists(img_s5):
        s18.shapes.add_picture(img_s5, Inches(0.8), Inches(1.8), width=Inches(5.6))
        cap5 = s18.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(5.6), Inches(0.5))
        p_c5 = cap5.text_frame.paragraphs[0]
        p_c5.text = "Figure 6.5: Multi-Hospital Diagnostic Reports Vault"
        p_c5.font.size = Pt(10)
        p_c5.font.bold = True
        p_c5.font.color.rgb = SLATE_TEXT
        p_c5.alignment = PP_ALIGN.CENTER

    if os.path.exists(img_s6):
        s18.shapes.add_picture(img_s6, Inches(6.8), Inches(1.8), width=Inches(5.7))
        cap6 = s18.shapes.add_textbox(Inches(6.8), Inches(6.5), Inches(5.7), Inches(0.5))
        p_c6 = cap6.text_frame.paragraphs[0]
        p_c6.text = "Figure 6.6: Sovereign Citizen DigiLocker & ABHA Identity Card"
        p_c6.font.size = Pt(10)
        p_c6.font.bold = True
        p_c6.font.color.rgb = SLATE_TEXT
        p_c6.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 19: PERFORMANCE & EXPERIMENTAL BENCHMARKS
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    add_header(s19, "Performance Benchmarking & Experimental Verification")

    # Metrics Boxes (4 cards)
    metrics = [
        ("< 15 ms", "Database Query Latency", "P95 response time for complete multi-hospital patient dossier retrieval across all entities."),
        ("< 50 ms", "AI Analysis Response", "Deterministic RAG pipeline, trajectory math, and synthesis output generation."),
        ("1.35 s", "Vite Production Build", "Full client compilation time with zero TypeScript compilation warnings or errors."),
        ("0 Errors", "Oxlint Quality Score", "Clean code verification across all 5 backend modules and frontend component hierarchy.")
    ]

    for idx, (stat, label, desc) in enumerate(metrics):
        x = Inches(0.8 + idx * 2.95)
        card = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.8), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_BLUE

        p_s = card.text_frame.paragraphs[0]
        p_s.text = stat
        p_s.font.size = Pt(28)
        p_s.font.bold = True
        p_s.font.color.rgb = PRIMARY_BLUE
        p_s.alignment = PP_ALIGN.CENTER

        p_l = card.text_frame.add_paragraph()
        p_l.text = label
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = DARK_NAVY
        p_l.alignment = PP_ALIGN.CENTER

        p_d = card.text_frame.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = SLATE_TEXT
        p_d.space_before = Pt(4)

    # Bottom Summary Table / Note
    card_bot = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.6))
    card_bot.fill.solid()
    card_bot.fill.fore_color.rgb = LIGHT_BG
    card_bot.line.color.rgb = BORDER_BLUE
    tf_b = card_bot.text_frame
    tf_b.margin_left = Inches(0.3)
    tf_b.margin_top = Inches(0.3)
    p_hb = tf_b.paragraphs[0]
    p_hb.text = "TEST HARNESS EXECUTION SUMMARY"
    p_hb.font.bold = True
    p_hb.font.size = Pt(13)
    p_hb.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_b, "Unit Test Matrix:", "6/6 automated route test cases passed (Auth, Dossier, Cure Certification, AI Analysis).", 11.5, 6)
    add_bullet_point(tf_b, "End-to-End Clinical Workflows:", "4/4 clinical scenarios validated (Metabolic onboarding, Dengue cure stamping, Lab vault ingestion, AI note copy).", 11.5, 6)
    add_bullet_point(tf_b, "Zero-Trust Security Test:", "Attempted horizontal privilege escalation from unauthorized doctor IDs blocked with 401/403.", 11.5, 6)

    # =========================================================================
    # SLIDE 20: CONCLUSION & FUTURE ROADMAP
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    add_header(s20, "Conclusion & Future Research Roadmap")

    # Left Card: Summary of Contributions
    card_c1 = s20.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    card_c1.fill.solid()
    card_c1.fill.fore_color.rgb = LIGHT_BG
    card_c1.line.color.rgb = BORDER_BLUE
    tf_c1 = card_c1.text_frame
    tf_c1.margin_left = Inches(0.3)
    tf_c1.margin_top = Inches(0.3)
    p_hc1 = tf_c1.paragraphs[0]
    p_hc1.text = "KEY PROJECT CONTRIBUTIONS"
    p_hc1.font.bold = True
    p_hc1.font.size = Pt(13)
    p_hc1.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_c1, "Unified Longitudinal Dossier:", "Unified multi-hospital records under an ABHA-compliant 14-digit citizen health identity.", 12, 10)
    add_bullet_point(tf_c1, "Active vs. Cured Continuity:", "Solved lifetime illness conflation through a certified cure state transition machine.", 12, 10)
    add_bullet_point(tf_c1, "Deterministic Clinical AI:", "Engineered zero-hallucination decision support with quantitative trajectory math and 1-click note copy.", 12, 10)
    add_bullet_point(tf_c1, "Zero-Trust Architecture:", "Enforced role-scoped workspaces in strict compliance with India's DPDPA 2023.", 12, 10)

    # Right Card: Future Scope
    card_c2 = s20.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3))
    card_c2.fill.solid()
    card_c2.fill.fore_color.rgb = CARD_BG
    card_c2.line.color.rgb = PRIMARY_BLUE
    tf_c2 = card_c2.text_frame
    tf_c2.margin_left = Inches(0.3)
    tf_c2.margin_top = Inches(0.3)
    p_hc2 = tf_c2.paragraphs[0]
    p_hc2.text = "FUTURE RESEARCH ROADMAP"
    p_hc2.font.bold = True
    p_hc2.font.size = Pt(13)
    p_hc2.font.color.rgb = PRIMARY_BLUE

    add_bullet_point(tf_c2, "Live NHA ABDM Gateway:", "Integrate official production sandbox certificates for real-time live hospital EHR synchronization.", 12, 10)
    add_bullet_point(tf_c2, "Wearable IoT Telemetry:", "Stream continuous glucose monitoring (CGM) and heart rate variability from consumer smartwatches.", 12, 10)
    add_bullet_point(tf_c2, "Vernacular Multilingual AI:", "Expand natural language copilot support to Hindi, Tamil, Telugu, and Bengali for rural health centers.", 12, 10)
    add_bullet_point(tf_c2, "Federated Clinical Research:", "Implement differential privacy for multi-hospital epidemiological research without exposing patient identities.", 12, 10)

    # =========================================================================
    # SLIDE 21: THANK YOU & Q&A
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    bg21 = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg21.fill.solid()
    bg21.fill.fore_color.rgb = DARK_NAVY
    bg21.line.fill.background()

    top21 = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    top21.fill.solid()
    top21.fill.fore_color.rgb = PRIMARY_BLUE
    top21.line.fill.background()

    tbox21 = s21.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.0))
    p_ty = tbox21.text_frame.paragraphs[0]
    p_ty.text = "THANK YOU!"
    p_ty.font.size = Pt(48)
    p_ty.font.bold = True
    p_ty.font.color.rgb = WHITE
    p_ty.alignment = PP_ALIGN.CENTER

    p_qa = tbox21.text_frame.add_paragraph()
    p_qa.text = "Questions & Answers (Viva-Voce Defense)"
    p_qa.font.size = Pt(22)
    p_qa.font.color.rgb = RGBColor(186, 230, 253)
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(10)

    p_repo = tbox21.text_frame.add_paragraph()
    p_repo.text = "Project Codebase: https://github.com/Samarth-27/MediSutra\nDepartment of Computer Science & Engineering | Academic Session 2025–2026"
    p_repo.font.size = Pt(13)
    p_repo.font.color.rgb = RGBColor(148, 163, 184)
    p_repo.alignment = PP_ALIGN.CENTER
    p_repo.space_before = Pt(30)

    output_ppt = "docs/MediSutra_Major_Project_Presentation.pptx"
    prs.save(output_ppt)
    print(f"Presentation successfully generated at: {output_ppt}")

if __name__ == '__main__':
    create_presentation()
